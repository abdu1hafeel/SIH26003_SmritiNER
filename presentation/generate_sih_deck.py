import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette: MDoNER North Eastern Care (Emerald, Deep Slate, Gold, Light Slate)
    C_DARK = RGBColor(15, 23, 42)       # Slate 900
    C_EMERALD = RGBColor(5, 122, 85)    # Emerald 700
    C_GOLD = RGBColor(217, 119, 6)      # Amber 600
    C_BG = RGBColor(248, 250, 252)      # Slate 50
    C_CARD_BG = RGBColor(255, 255, 255) # White
    C_TEXT = RGBColor(51, 65, 85)       # Slate 700
    C_MUTED = RGBColor(100, 116, 139)   # Slate 500
    C_WHITE = RGBColor(255, 255, 255)

    blank_layout = prs.slide_layouts[6]

    def add_header(slide, title_text, category_text="SIH26003 • Ministry of Development of North Eastern Region (MDoNER)"):
        # Header banner
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
        shape.fill.solid()
        shape.fill.fore_color.rgb = C_DARK
        shape.line.color.rgb = C_DARK

        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.4)
        tf.margin_top = Inches(0.18)

        p1 = tf.paragraphs[0]
        p1.text = category_text.upper()
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = C_GOLD

        p2 = tf.add_paragraph()
        p2.text = title_text
        p2.font.size = Pt(22)
        p2.font.bold = True
        p2.font.color.rgb = C_WHITE

    # -------------------------------------------------------------
    # SLIDE 1: TITLE SLIDE
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = C_DARK
    bg1.line.color.rgb = C_DARK

    title_box = slide1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(3.8))
    tf1 = title_box.text_frame
    tf1.word_wrap = True

    p_badge = tf1.paragraphs[0]
    p_badge.text = "SMART INDIA HACKATHON (SIH) 2026 • PROBLEM ID: SIH26003"
    p_badge.font.size = Pt(14)
    p_badge.font.bold = True
    p_badge.font.color.rgb = C_GOLD

    p_main = tf1.add_paragraph()
    p_main.text = "SmritiNER (স্মৃতি-উত্তৰ-পূব)"
    p_main.font.size = Pt(44)
    p_main.font.bold = True
    p_main.font.color.rgb = C_WHITE
    p_main.space_before = Pt(10)

    p_sub = tf1.add_paragraph()
    p_sub.text = "AI-Based Cognitive Gaming & Memory Assistance Platform for Elderly Dementia Patients in the North Eastern Region"
    p_sub.font.size = Pt(20)
    p_sub.font.color.rgb = RGBColor(203, 213, 225)
    p_sub.space_before = Pt(12)

    p_meta = tf1.add_paragraph()
    p_meta.text = "Ministry of Development of North Eastern Region (MDoNER) | Domain: MedTech / Software"
    p_meta.font.size = Pt(15)
    p_meta.font.bold = True
    p_meta.font.color.rgb = RGBColor(52, 211, 153)
    p_meta.space_before = Pt(20)

    # -------------------------------------------------------------
    # SLIDE 2: PROBLEM STATEMENT & GROUND REALITY IN NER
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "The Ground Reality & Geriatric Challenges in NER")

    cards_data_s2 = [
        ("Surging Elderly Dementia", "Mild Cognitive Impairment (MCI) and Alzheimer's cases are surging across NER states, projected to double in the next decade."),
        ("Geographical & Medical Isolation", "Hilly terrains and remote rural pockets severely restrict access to specialized geriatricians and memory care clinics."),
        ("Cultural & Linguistic Gap", "Mainstream memory apps feature foreign/urban concepts that alienate elderly patients in Assam, Manipur, Meghalaya, etc."),
        ("Connectivity & Power Bottlenecks", "Spotty cellular networks in remote valleys make cloud-dependent health applications completely unviable.")
    ]

    for i, (title, desc) in enumerate(cards_data_s2):
        col = i % 2
        row = i // 2
        left = Inches(0.8 + col * 5.95)
        top = Inches(1.8 + row * 2.5)
        
        box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.75), Inches(2.2))
        box.fill.solid()
        box.fill.fore_color.rgb = C_CARD_BG
        box.line.color.rgb = RGBColor(226, 232, 240)
        box.line.width = Pt(1.5)

        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.25)
        
        p = tf.paragraphs[0]
        p.text = f"{i+1}. {title}"
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = C_EMERALD

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(14)
        p2.font.color.rgb = C_TEXT
        p2.space_before = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 3: PROPOSED SOLUTION - SMRITINER
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "Proposed Solution: SmritiNER Platform Architecture")

    pillars = [
        ("4 Cognitive Domains", "Tailored to NER heritage: Memory (Bihu, Hornbill), Attention (Nature soundscapes), Routine Sequencing, and Textile Weave Patterns."),
        ("Adaptive AI Engine", "Mathematical latency & error monitoring. Provides gentle voice scaffolding when hesitating (>6s) to eliminate elder frustration."),
        ("Multi-Lingual Voice Guide", "Spoken regional prompts in Assamese (অসমীয়া), Meitei (মৈতৈলোন্), Bengali (বাংলা), Hindi, and English with comforting cadence."),
        ("Caregiver Telemetry & Geofence", "MoCA/MMSE clinical radar charts, early anomaly flags, medication compliance tracking, and GPS wandering safety alert.")
    ]

    for i, (title, desc) in enumerate(pillars):
        left = Inches(0.8 + i * 2.95)
        top = Inches(1.8)
        box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(2.85), Inches(5.0))
        box.fill.solid()
        box.fill.fore_color.rgb = C_CARD_BG
        box.line.color.rgb = C_EMERALD
        box.line.width = Pt(1.5)

        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = f"Pillar {i+1}"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = C_GOLD

        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(17)
        p2.font.bold = True
        p2.font.color.rgb = C_DARK
        p2.space_before = Pt(6)

        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(13)
        p3.font.color.rgb = C_TEXT
        p3.space_before = Pt(12)

    # -------------------------------------------------------------
    # SLIDE 4: SYSTEM FLOWCHART & TECHNICAL SPECIFICATIONS
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "System Architecture & Offline-First Flowchart")

    flow_box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(5.0))
    flow_box.fill.solid()
    flow_box.fill.fore_color.rgb = C_CARD_BG
    flow_box.line.color.rgb = RGBColor(203, 213, 225)

    tf_flow = flow_box.text_frame
    tf_flow.word_wrap = True
    tf_flow.margin_left = Inches(0.4)
    tf_flow.margin_top = Inches(0.3)

    p = tf_flow.paragraphs[0]
    p.text = "End-to-End System Data Flow:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = C_DARK

    steps = [
        ("Step 1: Patient Layer", "Ultra-accessible tablet interface (WCAG AAA high-contrast, >52px touch targets, regional speech guide)."),
        ("Step 2: AI Adaptive Loop", "Real-time touch latency & error evaluation: Delta L = alpha*(Target - Observed) + beta*(Accuracy - 0.75)."),
        ("Step 3: Auto-Scaffolding", "Dynamic audio hints & visual illumination activate if hesitation exceeds 6.0s to preserve dignity and eliminate stress."),
        ("Step 4: Offline Persistence", "100% offline-first execution via IndexedDB and Cache API; zero data loss in remote NER hilly zones."),
        ("Step 5: Caregiver & Doctor Tele-Health", "Auto-aggregates MoCA radar trends, triggers GPS wandering alarms, and exports clean clinical reports.")
    ]

    for title, desc in steps:
        p_step = tf_flow.add_paragraph()
        p_step.text = f"• {title}: {desc}"
        p_step.font.size = Pt(14)
        p_step.font.color.rgb = C_TEXT
        p_step.space_before = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 5: ELIMINATING AI SLOP & CLINICAL RIGOR
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "Eliminating 'AI Slop': Clinical Rigor & Authentic NER Design")

    comparison = [
        ("Typical Hackathon AI Slop", [
            "Generic buzzwords ('Deep AI magic') with no clinical grounding",
            "Western stock photos unrelated to North Eastern elders",
            "Stressful high-pressure countdowns triggering anxiety",
            "Cloud-only apps that break as soon as connection drops in hills",
            "Toy UI without doctor export or safety guardrails"
        ]),
        ("SmritiNER Clinical Engineering", [
            "Grounded in validated MoCA & MMSE cognitive assessment domains",
            "Authentic North Eastern assets: Bihu, Hornbill, Muga silk, Kaziranga",
            "Non-pharmacological sensory calming zone (bamboo flute & breathing)",
            "Resilient Offline-First PWA architecture with local IndexedDB",
            "PIN-protected caregiver portal, GPS wandering alerts & clinical PDF export"
        ])
    ]

    for i, (head, points) in enumerate(comparison):
        left = Inches(0.8 + i * 5.95)
        box = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.8), Inches(5.75), Inches(5.0))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(254, 242, 242) if i == 0 else RGBColor(240, 253, 244)
        box.line.color.rgb = RGBColor(239, 68, 68) if i == 0 else C_EMERALD
        box.line.width = Pt(1.5)

        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = head
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = RGBColor(185, 28, 28) if i == 0 else C_EMERALD

        for pt in points:
            p_pt = tf.add_paragraph()
            p_pt.text = f"• {pt}"
            p_pt.font.size = Pt(13)
            p_pt.font.color.rgb = C_DARK
            p_pt.space_before = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 6: FEASIBILITY, SENIOR ACCESSIBILITY & SECURITY
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "Clinical Feasibility, Accessibility & Safe-Zone Guard")

    feas_boxes = [
        ("WCAG AAA Accessibility", "Equipped with instant High-Contrast mode (black/amber) for cataract patients and scalable font typography up to 1.4x."),
        ("Procedural Folk Audio", "Web Audio API generates authentic bamboo flute pentatonic melodies in real-time, requiring 0kb external audio downloads."),
        ("Dementia Wandering Guard", "Simulated GPS perimeter (300m safe radius) that dispatches instant caregiver notifications if patient strays toward cliffs/rivers."),
        ("Doctor Consultation Ready", "One-click printable clinical PDF compiling 7-day reaction latency graphs, MoCA domain scores, and medication adherence.")
    ]

    for i, (title, desc) in enumerate(feas_boxes):
        col = i % 2
        row = i // 2
        left = Inches(0.8 + col * 5.95)
        top = Inches(1.8 + row * 2.5)

        box = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.75), Inches(2.2))
        box.fill.solid()
        box.fill.fore_color.rgb = C_CARD_BG
        box.line.color.rgb = RGBColor(203, 213, 225)
        box.line.width = Pt(1.5)

        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = C_EMERALD

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(14)
        p2.font.color.rgb = C_TEXT
        p2.space_before = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 7: IMPACT, SCALABILITY & MDONER ROADMAP
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "Impact, Scalability & MDoNER Rollout Roadmap")

    phases = [
        ("Phase 1: Hackathon Prototype (Months 0-3)", "Complete tablet PWA with 4 cognitive domains, offline IndexedDB engine, and caregiver portal tested with regional elders."),
        ("Phase 2: MDoNER Pilot (Months 3-6)", "Pilot deployment across Community Health Centers (CHCs) and Old Age Homes in Kamrup Metro, Imphal West, and East Khasi Hills."),
        ("Phase 3: ABDM & Tele-Health (Months 6-12)", "Direct integration with Ayushman Bharat Digital Mission (ABDM) and MDoNER telemedicine portals for remote specialist reviews."),
        ("Vision & Impact", "Dignified, culturally resonant cognitive health for over 2.5 million elderly citizens across North East India.")
    ]

    for i, (title, desc) in enumerate(phases):
        top = Inches(1.8 + i * 1.25)
        box = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top, Inches(11.733), Inches(1.05))
        box.fill.solid()
        box.fill.fore_color.rgb = C_CARD_BG
        box.line.color.rgb = C_EMERALD if i == 3 else RGBColor(226, 232, 240)
        box.line.width = Pt(1.5)

        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.12)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = C_EMERALD if i != 3 else C_GOLD

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(13)
        p2.font.color.rgb = C_TEXT
        p2.space_before = Pt(2)

    # Save Presentation
    output_path = os.path.join(os.path.dirname(__file__), "SIH26003_SmritiNER_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_deck()
